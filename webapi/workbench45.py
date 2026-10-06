"""Usable desktop workflows over the shared engine; not a second network stack."""
from __future__ import annotations
import base64
import csv
import importlib.util
import io
import ipaddress
import json
import os
import re
import secrets
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, replace
from datetime import datetime, timezone
from pathlib import Path
from types import SimpleNamespace
from typing import Literal
from fastapi import APIRouter, HTTPException, Request
from fastapi.responses import Response
from pydantic import BaseModel, Field, ConfigDict, field_validator
from webapi.runtime37 import connection, utcnow
from webapi.security37 import require_role, WRITE_RULES
from webapi.data37 import columns, ip_value, write_ping
from modules import auto_ip
from modules.icmp_probe import icmp_ping
from app_runtime import DATABASE_DIR, REPORT_DIR

router = APIRouter(prefix='/api/v45')
AUTO_LOCK = threading.RLock()
MAX_TARGETS = 2048
WRITE_RULES.extend([
    ('POST|PUT|DELETE', r'/api/v45/ping/targets(?:/[^/]+)?', ('Admin','Operator')),
    ('POST', r'/api/v45/ping/(start|stop|import-inventory|import)', ('Admin','Operator')),
    ('PUT|POST|DELETE', r'/api/v45/autoip/(targets(?:/[^/]+)?|options|import)', ('Admin',)),
    ('POST', r'/api/v45/autoip/(start|stop)', ('Admin','Operator')),
    ('POST', r'/api/v45/autoip/runs/[^/]+/report', ('Admin','Operator')),
    ('POST|PUT|DELETE', r'/api/v45/ipmac(?:/\d+)?', ('Admin','Operator')),
    ('POST', r'/api/v45/ipmac/import-scan', ('Admin',)),
    ('POST', r'/api/v45/ipmac/import-file', ('Admin','Operator')),
    ('POST', r'/api/v45/tasks', ('Admin','Operator')),
    ('POST', r'/api/v45/inventory/import', ('Admin',)),
    ('PUT', r'/api/v45/managed/\d+', ('Admin',)),
    ('POST', r'/api/v45/security/sync', ('Admin',)),
    ('POST', r'/api/v45/security/(events/\d+/close|assets/\d+/identity)', ('Admin','Operator')),
])


class StrictModel(BaseModel):
    model_config = ConfigDict(extra='forbid')


class TargetIn(StrictModel):
    ip: str = Field(min_length=1, max_length=45)
    name: str = Field(default='', max_length=200)
    profile: str = Field(default='', max_length=120)

    @field_validator('ip')
    @classmethod
    def unicast(cls, value):
        return auto_ip.validate_ip(value)


class TargetsIn(StrictModel):
    targets: list[TargetIn] = Field(max_length=MAX_TARGETS)
    mode: Literal['append','replace'] = 'append'
    confirmation: str = ''


class PingStart(StrictModel):
    authorized: bool = False
    repeat: bool = False
    interval: int = Field(default=15, ge=5, le=86400)
    timeout_ms: int = Field(default=1000, ge=100, le=10000)
    workers: int = Field(default=16, ge=1, le=32)


def ensure_tables():
    auto_ip.ensure_tables()
    with connection() as c:
        c.executescript('''
        CREATE TABLE IF NOT EXISTS web_ping_targets45(
          ip TEXT PRIMARY KEY,name TEXT NOT NULL DEFAULT '',status TEXT DEFAULT 'Unknown',
          response REAL,packet_loss REAL,observed_at TEXT,error TEXT DEFAULT '');
        CREATE TABLE IF NOT EXISTS web_ping_settings45(key TEXT PRIMARY KEY,value TEXT);
        ''')
        # Once only. An explicitly emptied list stays empty, including after restart.
        if not c.execute("SELECT 1 FROM web_ping_settings45 WHERE key='initialized'").fetchone():
            for row in c.execute('SELECT * FROM devices LIMIT ?', (MAX_TARGETS,)).fetchall():
                d = dict(row)
                try:
                    ip = auto_ip.validate_ip(ip_value(d))
                except ValueError:
                    continue
                c.execute('INSERT OR IGNORE INTO web_ping_targets45(ip,name) VALUES(?,?)', (ip,d.get('hostname') or ''))
            c.execute("INSERT INTO web_ping_settings45 VALUES('initialized','1')")


class PingCoordinator:
    """Bounded background ICMP, explicit start, immediate-progress snapshots.

    Mutation is rejected during a cycle. After stop completes no late result can
    resurrect a removed target. Repeating stops on server shutdown, not navigation.
    """
    def __init__(self):
        self.lock = threading.RLock()
        self.stop_event = threading.Event()
        self.thread = None
        self.state = {'status':'Ready','done':0,'total':0,'cycle':0,'running':False,
                      'next_run_at':None,'message':'','started_at':None,'finished_at':None}
        self.probe = icmp_ping

    def snapshot(self):
        with self.lock:
            return {**self.state,'running':bool(self.thread and self.thread.is_alive())}

    def start(self, options: PingStart, user):
        with self.lock:
            if self.thread and self.thread.is_alive():
                raise HTTPException(409,'PING_ALREADY_RUNNING')
            with connection() as c:
                targets = [dict(x) for x in c.execute('SELECT ip,name FROM web_ping_targets45 ORDER BY ip')]
            if not targets:
                raise HTTPException(400,'NO_PING_TARGETS: add IP addresses before starting')
            self.stop_event.clear()
            self.state = {'status':'Starting','done':0,'total':len(targets),'cycle':0,'running':True,
                          'next_run_at':None,'message':'','started_at':utcnow(),'finished_at':None,
                          'run_id':secrets.token_hex(16),'cycle_started_at':None,'owner':user['username'],'options':options.model_dump(exclude={'authorized'})}
            from webapi.operations47 import persist_ping
            persist_ping(self.state, user['id'])
            self.thread = threading.Thread(target=self._loop, args=(targets,options,user['id']),
                                           daemon=True,name='Web-Ping-Coordinator')
            self.thread.start()
            return dict(self.state)

    def stop(self, wait=False):
        with self.lock:
            self.stop_event.set()
            if self.thread and self.thread.is_alive():
                self.state['status'] = 'Stopping'
                self.state['next_run_at'] = None
        if wait and self.thread:
            self.thread.join(timeout=25)
        return self.snapshot()

    def assert_idle(self):
        if self.thread and self.thread.is_alive():
            raise HTTPException(409,'STOP_PING_FIRST: stop and wait for the current cycle before editing targets')

    def targets_changed(self, count):
        self.assert_idle()
        self.state.update(status="Ready",done=0,total=count,cycle=0,running=False,
                          next_run_at=None,started_at=None,finished_at=None,message="",run_id=None,cycle_started_at=None)
        with connection() as c:
            c.execute("UPDATE web_ping_signal47 SET bad=0,good=0")

    def _loop(self, targets, options, owner_id):
        try:
            while not self.stop_event.is_set():
                # Permissions are re-checked before each cycle (not just on submit).
                with connection() as c:
                    actor = c.execute('SELECT enabled,role FROM app_users WHERE id=?',(owner_id,)).fetchone()
                if not actor or not actor['enabled'] or actor['role'] not in ('Admin','Operator'):
                    raise RuntimeError('PERMISSION_REVOKED')
                with self.lock:
                    self.state.update(status='Running',done=0,cycle=self.state['cycle']+1,next_run_at=None,message='',cycle_started_at=utcnow())
                from webapi.operations47 import persist_ping, record_ping
                persist_ping(self.state)
                last_persist=time.monotonic()
                errors = 0
                with ThreadPoolExecutor(max_workers=options.workers,thread_name_prefix='Web-Ping') as pool:
                    def check(row):
                        if self.stop_event.is_set():
                            return None
                        try:
                            return self.probe(row['ip'],options.timeout_ms)
                        except Exception as exc:
                            return {'ip':row['ip'],'status':'Error','response':None,'packet_loss':None,
                                    'observed_at':utcnow(),'error':type(exc).__name__}
                    futures = {pool.submit(check,t):t for t in targets}
                    for f in as_completed(futures):
                        r = f.result()
                        if r is not None:
                            r.setdefault('observed_at',utcnow())
                            with connection() as c:
                                c.execute('''UPDATE web_ping_targets45 SET status=?,response=?,packet_loss=?,observed_at=?,error=? WHERE ip=?''',
                                          (r['status'],r.get('response'),r.get('packet_loss'),r['observed_at'],r.get('error',''),r['ip']))
                            write_ping(r)
                            record_ping(r, self.state['run_id'], self.state['cycle'])
                            errors += int(r['status'] in ('Error','Unknown'))
                            with self.lock:
                                self.state['done'] += 1
                            if time.monotonic()-last_persist>1:
                                persist_ping(self.state); last_persist=time.monotonic()
                        if self.stop_event.is_set():
                            for pending in futures:
                                pending.cancel()
                            break
                if not options.repeat or self.stop_event.is_set():
                    with self.lock:
                        self.state['status'] = 'Stopped' if self.stop_event.is_set() else 'CompletedWithErrors' if errors else 'Completed'
                    break
                with self.lock:
                    self.state.update(status='Waiting',next_run_at=datetime.fromtimestamp(time.time()+options.interval,timezone.utc).isoformat())
                persist_ping(self.state)
                self.stop_event.wait(options.interval)
        except Exception as exc:
            with self.lock:
                self.state.update(status='Failed',message=str(exc)[:300])
        finally:
            with self.lock:
                if self.stop_event.is_set() and self.state['status'] != 'Failed':
                    self.state['status'] = 'Stopped'
                self.state.update(running=False,next_run_at=None,finished_at=utcnow())
                from webapi.operations47 import persist_ping
                persist_ping(self.state)


ping = PingCoordinator()


@router.get('/ping')
def ping_status(request:Request):
    require_role(request)
    with connection() as c:
        targets = [dict(x) for x in c.execute('SELECT * FROM web_ping_targets45 ORDER BY ip')]
    from webapi.data37 import freshness
    for t in targets:
        t['data_state'] = freshness(t.get('observed_at'))
    snapshot=ping.snapshot()
    if snapshot['status']=='Ready': snapshot['total']=len(targets)
    from webapi.operations47 import enrich_ping
    current_summary=enrich_ping(targets,snapshot)
    return {'state':snapshot,'targets':targets,'current_summary':current_summary,
            'summary':{s:sum(t['status']==s for t in targets) for s in ('Online','Offline','Unknown','Error')}}


@router.post('/ping/start',status_code=202)
def ping_start(x:PingStart, request:Request):
    user=require_role(request,'Admin','Operator')
    if not x.authorized:
        raise HTTPException(400,'AUTHORIZATION_REQUIRED')
    return {'success':True,'state':ping.start(x,user)}


@router.post('/ping/stop')
def ping_stop(request:Request):
    require_role(request,'Admin','Operator')
    return {'success':True,'state':ping.stop()}


@router.put('/ping/targets')
def ping_targets(x:TargetsIn,request:Request):
    require_role(request,'Admin','Operator')
    with ping.lock:
        ping.assert_idle()
        with connection() as c:
            c.execute('BEGIN IMMEDIATE')
            if x.mode == 'replace':
                if x.confirmation != 'REPLACE':
                    raise HTTPException(400,'Type REPLACE to replace the target list')
                c.execute('DELETE FROM web_ping_targets45')
            for t in x.targets:
                c.execute('''INSERT INTO web_ping_targets45(ip,name) VALUES(?,?)
                             ON CONFLICT(ip) DO UPDATE SET name=excluded.name''',(t.ip,t.name))
            count=c.execute('SELECT COUNT(*) FROM web_ping_targets45').fetchone()[0]
            if count>MAX_TARGETS:
                raise HTTPException(400,'Maximum 2048 targets')
        ping.targets_changed(count)
    return {'success':True,'count':count}


@router.delete('/ping/targets/{ip}')
def ping_delete(ip:str,request:Request):
    require_role(request,'Admin','Operator')
    with ping.lock:
        ping.assert_idle()
        with connection() as c:
            if not c.execute('DELETE FROM web_ping_targets45 WHERE ip=?',(ip,)).rowcount:
                raise HTTPException(404,'Target not found')
            count=c.execute('SELECT COUNT(*) FROM web_ping_targets45').fetchone()[0]
        ping.targets_changed(count)
    return {'success':True}


@router.post('/ping/import-inventory')
def ping_inventory(request:Request):
    require_role(request,'Admin','Operator')
    with connection() as c:
        records=[dict(x) for x in c.execute('SELECT * FROM devices LIMIT ?',(MAX_TARGETS,))]
    targets=[]
    for d in records:
        try:
            targets.append(TargetIn(ip=ip_value(d),name=d.get('hostname') or ''))
        except ValueError:
            continue
    return ping_targets(TargetsIn(targets=targets),request)


def auto_idle():
    from webapi import main
    if main.AUTO_ENGINE and main.AUTO_ENGINE.running:
        raise HTTPException(409,'STOP_AUTOIP_FIRST: wait for the active run to stop before changing targets/options')


@router.put('/autoip/targets')
def auto_targets(x:TargetsIn,request:Request):
    require_role(request,'Admin')
    with AUTO_LOCK:
        auto_idle()
        old={t.ip:t for t in auto_ip.load_targets()} if x.mode=='append' else {}
        if x.mode=='replace' and x.confirmation!='REPLACE':
            raise HTTPException(400,'Type REPLACE to replace the target list')
        known={p['name'] for p in auto_ip.profiles()}
        for t in x.targets:
            if t.profile and t.profile not in known:
                raise HTTPException(400,'UNKNOWN_PROFILE: '+t.profile)
            old[t.ip]=auto_ip.Target(t.ip,t.name,t.profile)
        if len(old)>MAX_TARGETS:
            raise HTTPException(400,'Maximum 2048 targets')
        auto_ip.save_targets(list(old.values()))
    return {'success':True,'count':len(old)}


@router.delete('/autoip/targets/{ip}')
def auto_delete(ip:str,request:Request):
    require_role(request,'Admin')
    with AUTO_LOCK:
        auto_idle()
        with connection() as c:
            if not c.execute('DELETE FROM autoip_targets WHERE ip=?',(ip,)).rowcount:
                raise HTTPException(404,'Target not found')
    return {'success':True}


class AutoOptions(StrictModel):
    default_profile:str=Field(default=auto_ip.DEFAULT_PROFILE,min_length=1,max_length=120)
    ping:bool=True; tcp:bool=True; snmp:bool=True
    resources:bool=True; interfaces:bool=True; topology:bool=True
    backup:bool=False; alerts:bool=True; notifications:bool=False; report:bool=True
    repeat:bool=False
    interval:int=Field(default=300,ge=30,le=86400)
    workers:int=Field(default=4,ge=1,le=16)
    timeout:float=Field(default=2,ge=.5,le=10)
    retries:int=Field(default=1,ge=0,le=2)
    max_interfaces:int=Field(default=128,ge=1,le=512)
    device_budget:int=Field(default=120,ge=30,le=600)
    ports:str=Field(default='22,80,443',max_length=200)
    # Web starts explicitly. Opening a browser must not send traffic or mail.
    auto_start:bool=False
    email_report_after_run:bool=False
    report_email_to:str=Field(default='',max_length=254)

    def options(self):
        try:
            return auto_ip.Options(**self.model_dump()).validate()
        except ValueError as exc:
            raise HTTPException(400,str(exc)) from exc


@router.put('/autoip/options')
def auto_options(x:AutoOptions,request:Request):
    require_role(request,'Admin')
    with AUTO_LOCK:
        auto_idle()
        if x.default_profile not in {p['name'] for p in auto_ip.profiles()}:
            raise HTTPException(400,'UNKNOWN_DEFAULT_PROFILE')
        auto_ip.save_options(x.options())
    return {'success':True,'options':x.model_dump()}


class AutoStart(StrictModel):
    authorized:bool=False
    options:AutoOptions
    confirm_notifications:bool=False


@router.post('/autoip/start',status_code=202)
def auto_start(x:AutoStart,request:Request):
    from webapi import main
    user=require_role(request,'Admin','Operator')
    if not x.authorized:
        raise HTTPException(400,'AUTHORIZATION_REQUIRED')
    if os.environ.get('NA_ENABLE_AUTOIP') != '1':
        raise HTTPException(409,'AUTOIP_DISABLED: launch using START_WEB.bat or set NA_ENABLE_AUTOIP=1')
    if x.options.notifications or x.options.email_report_after_run:
        if user['role']!='Admin' or not x.confirm_notifications:
            raise HTTPException(403,'Admin must explicitly confirm outbound notifications')
    with AUTO_LOCK:
        auto_idle()
        targets=auto_ip.load_targets()
        if not targets:
            raise HTTPException(400,'NO_AUTOIP_TARGETS')
        if main.AUTO_ENGINE is None:
            main.AUTO_ENGINE=auto_ip.AutomationEngine()
        try:
            main.AUTO_ENGINE.start(targets,x.options.options(),user,authorized=True)
        except (ValueError,RuntimeError,PermissionError) as exc:
            raise HTTPException(409,str(exc)) from exc
        return {'success':True,'state':main.AUTO_ENGINE.snapshot(),'targets':len(targets)}


@router.post('/autoip/stop')
def auto_stop(request:Request):
    from webapi import main
    require_role(request,'Admin','Operator')
    if main.AUTO_ENGINE:
        main.AUTO_ENGINE.stop()
    return {'success':True,'state':main.api_autoip_status()}


@router.get('/autoip/runs/{run_id}')
def auto_run(run_id:str,request:Request,after:int=0,limit:int=500):
    require_role(request,'Admin','Operator')
    with connection() as c:
        row=c.execute('SELECT id,started_at,finished_at,status,username,total,cycle FROM autoip_runs WHERE id=?',(run_id,)).fetchone()
        if not row:
            raise HTTPException(404,'Run not found')
        counts={r['status']:r['n'] for r in c.execute('SELECT status,COUNT(*) n FROM autoip_steps WHERE run_id=? GROUP BY status',(run_id,))}
        rows=[dict(r) for r in c.execute('SELECT * FROM autoip_steps WHERE run_id=? AND id>? ORDER BY id LIMIT ?',
                                      (run_id,max(0,after),max(1,min(limit,2000))))]
    for r in rows:
        try:
            r['data']=json.loads(r['data'])
        except (ValueError,TypeError):
            r['data']={}
    return {'run':dict(row),'summary':counts,'steps':rows,'next_after':rows[-1]['id'] if rows else after}


@router.post('/autoip/runs/{run_id}/report')
def auto_report(run_id:str,request:Request):
    user=require_role(request,'Admin','Operator')
    with connection() as c:
        r=c.execute('SELECT id FROM autoip_runs WHERE id=?',(run_id,)).fetchone()
    if not r:
        raise HTTPException(404,'Run not found')
    from webapi.reports37 import register
    target=Path(REPORT_DIR)/('autoip_'+secrets.token_hex(10)+'.xlsx')
    auto_ip.export_report(run_id,target)
    return register(target,user['id'],'report')


class ImportIn(StrictModel):
    filename:str=Field(max_length=200)
    content_base64:str=Field(max_length=12*1024*1024)
    mode:Literal['append','replace']='append'
    confirmation:str=''


def parse_import(x:ImportIn):
    try:
        data=base64.b64decode(x.content_base64,validate=True)
    except ValueError as exc:
        raise HTTPException(400,'INVALID_FILE_ENCODING') from exc
    if len(data)>8*1024*1024:
        raise HTTPException(413,'File limit is 8 MB')
    suffix=Path(x.filename).suffix.lower()
    try:
        if suffix=='.xlsx':
            with tempfile.TemporaryDirectory(prefix='na-import-') as d:
                path=Path(d)/'targets.xlsx';path.write_bytes(data)
                targets,duplicates=auto_ip.parse_excel(path)
        elif suffix in ('.csv','.txt'):
            text=data.decode('utf-8-sig')
            if suffix=='.csv':
                rows=list(csv.DictReader(io.StringIO(text)))
                normalized=[]
                aliases={'ip':'ip','ipaddress':'ip','diachiip':'ip','name':'name','hostname':'name','tenthietbi':'name','profile':'profile','hoso':'profile'}
                for row in rows:
                    normalized.append({aliases.get(auto_ip.normalize_header(k),k):v for k,v in row.items() if k})
                targets,duplicates=auto_ip.parse_rows(normalized)
            else:
                targets,duplicates=auto_ip.parse_text(text)
        else:
            raise ValueError('Supported files: .xlsx, .csv, .txt')
        return targets,duplicates
    except (ValueError,UnicodeError,KeyError) as exc:
        raise HTTPException(400,str(exc)) from exc


@router.post('/autoip/import')
def auto_import(x:ImportIn,request:Request):
    require_role(request,'Admin')
    targets,duplicates=parse_import(x)
    result=auto_targets(TargetsIn(targets=[TargetIn(**asdict(t)) for t in targets],mode=x.mode,confirmation=x.confirmation),request)
    return {**result,'duplicates_skipped':duplicates}


@router.post('/inventory/import')
def inventory_import(x:ImportIn,request:Request):
    require_role(request,'Admin')
    if x.mode!='append':
        raise HTTPException(400,'Inventory import is append-only; existing devices and credentials are preserved')
    targets,duplicates=parse_import(x)
    from webapi import main
    created=[];existing=[]
    with connection() as c:
        known={ip_value(dict(r)) for r in c.execute('SELECT * FROM devices')}
    for t in targets:
        if t.ip in known:
            existing.append(t.ip);continue
        created.append(main.add_device(main.DeviceIn(ip=t.ip,hostname=t.name)))
        known.add(t.ip)
    return {'success':True,'created':len(created),'existing':len(existing),'duplicates_skipped':duplicates}


class IPMacIn(StrictModel):
    ip:str=Field(max_length=45)
    mac:str=Field(default='',max_length=32)
    hostname:str=Field(default='',max_length=200)
    note:str=Field(default='',max_length=1000)

    @field_validator('ip')
    @classmethod
    def valid_ip(cls,v):
        return auto_ip.validate_ip(v)

    @field_validator('mac')
    @classmethod
    def valid_mac(cls,v):
        v=v.strip().upper().replace('-',':')
        if v and not re.fullmatch(r'(?:[0-9A-F]{2}:){5}[0-9A-F]{2}',v):
            raise ValueError('MAC format: AA:BB:CC:DD:EE:FF')
        return v


@router.get('/ipmac')
def ipmac_get(request:Request):
    require_role(request)
    with connection() as c:
        result=[dict(x) for x in c.execute('SELECT * FROM ip_mac_inventory ORDER BY ip')]
    counts={}
    for r in result:
        if r.get('mac'):
            key=r['mac'].replace('-',':').upper();counts[key]=counts.get(key,0)+1
    now=datetime.now(timezone.utc)
    for r in result:
        r['conflict']='MAC multi-IP' if counts.get((r.get('mac') or '').replace('-',':').upper(),0)>1 else ''
        r['observed_status']=r.get('status') or 'Unknown'
        raw=r.get('last_seen')
        age=None
        if raw:
            text=str(raw).strip().replace('Z','+00:00')
            try:
                parsed=datetime.fromisoformat(text)
            except ValueError:
                try: parsed=datetime.strptime(text[:19],'%Y-%m-%d %H:%M:%S')
                except ValueError: parsed=None
            if parsed is not None:
                if parsed.tzinfo is None: parsed=parsed.replace(tzinfo=timezone.utc)
                age=max(0,int((now-parsed.astimezone(timezone.utc)).total_seconds()))
        r['age_seconds']=age
        if age is None:
            r['freshness']='UNKNOWN';r['effective_status']='UNKNOWN'
        elif age<=180:
            r['freshness']='FRESH';r['effective_status']=r['observed_status']
        else:
            r['freshness']='STALE';r['effective_status']='STALE'
    return result


@router.post('/ipmac')
def ipmac_add(x:IPMacIn,request:Request):
    require_role(request,'Admin','Operator')
    with connection() as c:
        if c.execute('SELECT 1 FROM ip_mac_inventory WHERE ip=?',(x.ip,)).fetchone():
            raise HTTPException(409,'IP_ALREADY_EXISTS')
        c.execute('INSERT INTO ip_mac_inventory(ip,mac,hostname,note,status,first_seen) VALUES(?,?,?,?,?,?)',
                  (x.ip,x.mac,x.hostname,x.note,'Unknown',utcnow()))
    return {'success':True}


@router.put('/ipmac/{row_id}')
def ipmac_edit(row_id:int,x:IPMacIn,request:Request):
    require_role(request,'Admin','Operator')
    with connection() as c:
        if c.execute('SELECT 1 FROM ip_mac_inventory WHERE ip=? AND id<>?',(x.ip,row_id)).fetchone():
            raise HTTPException(409,'IP_ALREADY_EXISTS')
        if not c.execute('UPDATE ip_mac_inventory SET ip=?,mac=?,hostname=?,note=? WHERE id=?',
                         (x.ip,x.mac,x.hostname,x.note,row_id)).rowcount:
            raise HTTPException(404,'Record not found')
    return {'success':True}


@router.delete('/ipmac/{row_id}')
def ipmac_delete(row_id:int,request:Request):
    require_role(request,'Admin','Operator')
    with connection() as c:
        if not c.execute('DELETE FROM ip_mac_inventory WHERE id=?',(row_id,)).rowcount:
            raise HTTPException(404,'Record not found')
    return {'success':True}


class IPMacDeleteMany(StrictModel):
    ids:list[int]=Field(min_length=1,max_length=5000)
    confirmation:Literal['DELETE_SELECTED']


class IPMacDeleteAll(StrictModel):
    confirmation:Literal['DELETE_ALL_IPMAC']


def _delete_ipmac_ids(c,ids):
    unique=sorted({int(x) for x in ids if int(x)>0})
    deleted=0
    for pos in range(0,len(unique),400):
        chunk=unique[pos:pos+400]
        q=','.join('?' for _ in chunk)
        deleted+=c.execute(f'DELETE FROM ip_mac_inventory WHERE id IN ({q})',chunk).rowcount
    return deleted


@router.post('/ipmac/delete-many')
def ipmac_delete_many(x:IPMacDeleteMany,request:Request):
    require_role(request,'Admin','Operator')
    with connection() as c:
        deleted=_delete_ipmac_ids(c,x.ids)
    return {'success':True,'deleted':deleted}


@router.post('/ipmac/delete-all')
def ipmac_delete_all(x:IPMacDeleteAll,request:Request):
    require_role(request,'Admin')
    with connection() as c:
        deleted=c.execute('DELETE FROM ip_mac_inventory').rowcount
    return {'success':True,'deleted':deleted}


class ManagedEdit(StrictModel):
    name:str=Field(min_length=1,max_length=200)
    vendor:str=Field(default='',max_length=120)
    model:str=Field(default='',max_length=120)
    device_type:str=Field(default='',max_length=100)
    location:str=Field(default='',max_length=200)
    note:str=Field(default='',max_length=1000)


@router.put('/managed/{device_id}')
def managed_edit(device_id:int,x:ManagedEdit,request:Request):
    require_role(request,'Admin')
    with connection() as c:
        cols=columns(c,'network_devices');values={k:v for k,v in x.model_dump().items() if k in cols}
        if 'device_name' in cols:
            values['device_name']=x.name
        values['updated_at']=utcnow()
        if not c.execute('UPDATE network_devices SET '+','.join(k+'=?' for k in values)+' WHERE id=?',
                         list(values.values())+[device_id]).rowcount:
            raise HTTPException(404,'Managed device not found')
    return {'success':True}


@router.get('/remote/{device_id}/rdp')
def remote_rdp(device_id:int,request:Request):
    require_role(request,'Admin','Operator')
    with connection() as c:
        r=c.execute('SELECT * FROM network_devices WHERE id=?',(device_id,)).fetchone()
    if not r:
        raise HTTPException(404,'Managed device not found')
    ip=str(ipaddress.ip_address(ip_value(dict(r))))
    text=f'full address:s:{ip}:3389\r\nprompt for credentials:i:1\r\nauthentication level:i:2\r\nredirectclipboard:i:0\r\n'
    return Response(text,media_type='application/x-rdp',headers={'Content-Disposition':f'attachment; filename="device-{device_id}.rdp"'})


EXTENDED_ROLES={'RESOURCE_POLL':('Admin','Operator'),'INTERFACE_POLL':('Admin','Operator'),
                'REMOTE_CHECK':('Admin','Operator'),'NOTIFICATION_TEST':('Admin',),
                'SECURITY_TCP':('Admin','Operator'),'SECURITY_TLS':('Admin','Operator')}


class TaskIn(StrictModel):
    operation:Literal['RESOURCE_POLL','INTERFACE_POLL','REMOTE_CHECK','NOTIFICATION_TEST','SECURITY_TCP','SECURITY_TLS']
    device_ids:list[int]=Field(default_factory=list,max_length=128)
    asset_ids:list[int]=Field(default_factory=list,max_length=32)
    ifindices:list[int]=Field(default_factory=lambda:[1],max_length=32)
    service:Literal['SSH','TELNET','HTTP','HTTPS','RDP']='SSH'
    channel:Literal['Telegram','Email']='Email'
    authorized:bool=False


@router.post('/tasks',status_code=202)
def extended_submit(x:TaskIn,request:Request):
    from webapi.jobs37 import engine
    user=require_role(request,*EXTENDED_ROLES[x.operation])
    if not x.authorized:
        raise HTTPException(400,'AUTHORIZATION_REQUIRED')
    ids=list(dict.fromkeys(x.device_ids))
    if x.operation.startswith('SECURITY_'):
        if not x.asset_ids or min(x.asset_ids)<1:
            raise HTTPException(400,'Select registered security assets')
        with connection() as c:
            found={r['id'] for r in c.execute('SELECT id FROM security_assets WHERE id IN ('+','.join('?' for _ in x.asset_ids)+')',x.asset_ids)}
        if found!=set(x.asset_ids):
            raise HTTPException(404,'Unregistered security asset')
    elif x.operation!='NOTIFICATION_TEST':
        if not ids or min(ids)<1:
            raise HTTPException(400,'Select registered managed devices')
        with connection() as c:
            found={r['id'] for r in c.execute('SELECT id FROM network_devices WHERE id IN ('+','.join('?' for _ in ids)+')',ids)}
        if found!=set(ids):
            raise HTTPException(404,'Unregistered managed device ID')
    if x.operation=='INTERFACE_POLL' and (not x.ifindices or min(x.ifindices)<1 or max(x.ifindices)>2147483647):
        raise HTTPException(400,'Invalid IfIndex')
    return engine.submit(x.operation,{**x.model_dump(exclude={'authorized'}),'device_ids':ids},user)


def execute_extended(operation,payload,user,progress,cancelled):
    from webapi import parity44 as p
    req=SimpleNamespace(state=SimpleNamespace(user=user))
    if operation=='NOTIFICATION_TEST':
        # Pydantic model names are kept at the existing public endpoint boundary.
        return p.test_notification(p.NotificationTestIn(channel=payload['channel'],authorized=True),req)
    if operation.startswith('SECURITY_'):
        from modules import enterprise_security as security
        ids=payload['asset_ids'];results=[];progress(total=len(ids),done=0)
        for asset_id in ids:
            if cancelled(): break
            with connection() as c:
                row=c.execute('SELECT ip FROM security_assets WHERE id=?',(asset_id,)).fetchone()
            try:
                if not row: raise ValueError('Asset removed before execution')
                result=security.tcp_probe(row['ip']) if operation=='SECURITY_TCP' else security.tls_certificate_check(row['ip'])
                result['interpretation']='Observed open ports; not a security certification' if operation=='SECURITY_TCP' else 'TLS handshake and fingerprint only; certificate trust/hostname validation not performed'
                results.append({'asset_id':asset_id,'ip':row['ip'],'success':True,'result':result})
            except Exception as exc:
                results.append({'asset_id':asset_id,'success':False,'detail':str(exc)[:300]})
            progress(done=len(results))
        return results
    results=[];ids=payload['device_ids'];progress(total=len(ids),done=0)
    for did in ids:
        if cancelled():
            break
        try:
            if operation=='RESOURCE_POLL':
                r=p.poll_resource(did,p.ResourcePollIn(authorized=True),req)
            elif operation=='INTERFACE_POLL':
                r=p.poll_interfaces(did,p.InterfacePollIn(authorized=True,ifindices=payload['ifindices']),req)
            else:
                r=p.remote_check(did,p.RemoteCheckIn(authorized=True,service=payload['service']),req)
            results.append({'device_id':did,'success':r.get('success',False),'result':r})
        except HTTPException as exc:
            results.append({'device_id':did,'success':False,'detail':str(exc.detail)})
        except Exception as exc:
            results.append({'device_id':did,'success':False,'detail':type(exc).__name__})
        progress(done=len(results))
    return results


@router.get('/capabilities')
def capabilities(request:Request):
    require_role(request)
    libraries={name:bool(importlib.util.find_spec(name)) for name in ('paramiko','pysnmp','openpyxl')}
    import shutil
    return {'version':'4.5.0','platform':os.name,'ping_available':bool(shutil.which('ping')),
            'libraries':libraries,'max_targets':MAX_TARGETS,
            'autoip_enabled':os.environ.get('NA_ENABLE_AUTOIP')=='1',
            'daily_audit_available':os.name=='nt',
            'note':'Availability of code/dependencies is not proof of device connectivity.'}


@router.post('/ping/import')
def ping_import(x:ImportIn,request:Request):
    require_role(request,'Admin','Operator')
    targets,duplicates=parse_import(x)
    result=ping_targets(TargetsIn(targets=[TargetIn(**asdict(t)) for t in targets],mode=x.mode,confirmation=x.confirmation),request)
    return {**result,'duplicates_skipped':duplicates}


def _norm_ipmac_header(value):
    value=auto_ip.normalize_header(str(value or ''))
    aliases={
        'ip':'ip','ipaddress':'ip','diachiip':'ip','diachi':'ip',
        'mac':'mac','macaddress':'mac','diachimac':'mac',
        'hostname':'hostname','host':'hostname','name':'hostname','tenthietbi':'hostname','thietbi':'hostname',
        'note':'note','notes':'note','ghichu':'note','mota':'note','description':'note',
    }
    return aliases.get(value, value)


def _validate_ipmac_row(row, line_no):
    normalized={_norm_ipmac_header(k): ('' if v is None else str(v).strip()) for k,v in (row or {}).items() if k is not None}
    ip=normalized.get('ip','').strip()
    if not ip:
        raise ValueError(f'Dòng {line_no}: thiếu cột IP')
    try:
        ip=auto_ip.validate_ip(ip)
    except Exception as exc:
        raise ValueError(f'Dòng {line_no}: IP không hợp lệ: {ip}') from exc
    mac=normalized.get('mac','').strip().upper().replace('-',':')
    if mac and not re.fullmatch(r'(?:[0-9A-F]{2}:){5}[0-9A-F]{2}',mac):
        raise ValueError(f'Dòng {line_no}: MAC không hợp lệ: {mac}')
    hostname=normalized.get('hostname','').strip()[:200]
    note=normalized.get('note','').strip()[:1000]
    return {'ip':ip,'mac':mac,'hostname':hostname,'note':note}


def parse_ipmac_import(x:ImportIn):
    try:
        data=base64.b64decode(x.content_base64,validate=True)
    except ValueError as exc:
        raise HTTPException(400,'INVALID_FILE_ENCODING') from exc
    if len(data)>8*1024*1024:
        raise HTTPException(413,'File limit is 8 MB')
    suffix=Path(x.filename).suffix.lower()
    raw_rows=[]
    try:
        if suffix=='.xlsx':
            if not importlib.util.find_spec('openpyxl'):
                raise ValueError('Thiếu thư viện openpyxl. Hãy chạy INSTALL_WEB.bat rồi thử lại.')
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            ws=wb.active
            rows=ws.iter_rows(values_only=True)
            header=next(rows,None)
            if not header:
                raise ValueError('File Excel không có dữ liệu')
            keys=[str(v or '').strip() for v in header]
            for idx,values in enumerate(rows,start=2):
                if not any(v is not None and str(v).strip() for v in values):
                    continue
                raw_rows.append((idx,{keys[i]:values[i] if i<len(values) else '' for i in range(len(keys))}))
        elif suffix=='.csv':
            text=data.decode('utf-8-sig')
            try:
                dialect=csv.Sniffer().sniff(text[:4096],delimiters=',;\t')
            except csv.Error:
                dialect=csv.excel
            reader=csv.DictReader(io.StringIO(text),dialect=dialect)
            if not reader.fieldnames:
                raise ValueError('CSV phải có hàng tiêu đề')
            for idx,row in enumerate(reader,start=2):
                if not any(str(v or '').strip() for v in row.values()):
                    continue
                raw_rows.append((idx,row))
        elif suffix=='.txt':
            text=data.decode('utf-8-sig')
            lines=[line.strip() for line in text.splitlines() if line.strip() and not line.lstrip().startswith('#')]
            for idx,line in enumerate(lines,start=1):
                parts=[p.strip() for p in re.split(r'[;,\t]',line,maxsplit=3)]
                raw_rows.append((idx,{'IP':parts[0] if parts else '', 'MAC':parts[1] if len(parts)>1 else '', 'Hostname':parts[2] if len(parts)>2 else '', 'Note':parts[3] if len(parts)>3 else ''}))
        else:
            raise ValueError('Chỉ hỗ trợ file .xlsx, .csv hoặc .txt')
        if not raw_rows:
            raise ValueError('File không có thiết bị để nhập')
        if len(raw_rows)>5000:
            raise ValueError('Tối đa 5000 dòng thiết bị mỗi lần nhập')
        result=[];seen=set();duplicates=0
        for line_no,row in raw_rows:
            item=_validate_ipmac_row(row,line_no)
            if item['ip'] in seen:
                duplicates+=1
                continue
            seen.add(item['ip']);result.append(item)
        return result,duplicates
    except (ValueError,UnicodeError,KeyError) as exc:
        raise HTTPException(400,str(exc)) from exc


@router.post('/ipmac/import-file')
def ipmac_import_file(x:ImportIn,request:Request):
    require_role(request,'Admin','Operator')
    if x.mode!='append':
        raise HTTPException(400,'IP/MAC file import is append-only; existing records are preserved')
    items,duplicates=parse_ipmac_import(x)
    created=0;existing=[]
    now=utcnow()
    with connection() as c:
        known={str(r['ip']).strip() for r in c.execute('SELECT ip FROM ip_mac_inventory').fetchall()}
        for item in items:
            if item['ip'] in known:
                existing.append(item['ip'])
                continue
            c.execute("INSERT INTO ip_mac_inventory(ip,mac,hostname,note,status,first_seen) VALUES(?,?,?,?,?,?)",
                      (item['ip'],item['mac'],item['hostname'],item['note'],'Unknown',now))
            known.add(item['ip']);created+=1
    return {'success':True,'created':created,'existing':len(existing),'existing_ips':existing[:50],
            'duplicates_skipped':duplicates,'total_rows':len(items)+duplicates}


@router.post('/ipmac/import-scan')
def ipmac_import_scan(request:Request):
    require_role(request,'Admin')
    with connection() as c:
        latest=c.execute('SELECT scan_key FROM web_scan_results ORDER BY id DESC LIMIT 1').fetchone()
        if not latest:
            raise HTTPException(400,'NO_SCAN_RESULTS: run network discovery first')
        rows=[dict(r) for r in c.execute("SELECT * FROM web_scan_results WHERE scan_key=? AND lower(status) IN ('online','activearp','activelan')",(latest['scan_key'],))]
        for r in rows:
            address=auto_ip.validate_ip(r['ip'])
            c.execute('''INSERT INTO ip_mac_inventory(ip,mac,hostname,status,note,first_seen,last_seen)
                         VALUES(?,?,?,?,'',?,?) ON CONFLICT(ip) DO UPDATE SET
                         mac=CASE WHEN excluded.mac<>'' THEN excluded.mac ELSE ip_mac_inventory.mac END,
                         hostname=CASE WHEN excluded.hostname<>'' THEN excluded.hostname ELSE ip_mac_inventory.hostname END,
                         status=excluded.status,last_seen=excluded.last_seen''',
                      (address,r.get('mac') or '',r.get('hostname') or '',('Active' if str(r.get('status') or '').lower() in ('activearp','activelan') else 'Online'),r['created_at'],r['created_at']))
    return {'success':True,'imported':len(rows),'scan_key':latest['scan_key']}


@router.post('/security/sync')
def security_sync(request:Request):
    require_role(request,'Admin')
    from modules.enterprise_security import sync_registered_assets
    return {'success':True,'count':sync_registered_assets()}


@router.post('/security/events/{event_id}/close')
def security_close(event_id:int,request:Request):
    require_role(request,'Admin','Operator')
    from modules.enterprise_security import close_event, ensure_enterprise_security_tables
    ensure_enterprise_security_tables()
    with connection() as c:
        if not c.execute('SELECT 1 FROM security_events WHERE id=?',(event_id,)).fetchone():
            raise HTTPException(404,'Event not found')
    close_event(event_id)
    return {'success':True}


class IdentityIn(StrictModel):
    observed_mac:str=Field(default='',max_length=32)
    observed_hostname:str=Field(default='',max_length=200)
    authorized:bool=False

    @field_validator('observed_mac')
    @classmethod
    def mac(cls,value):
        return IPMacIn.valid_mac(value)


@router.post('/security/assets/{asset_id}/identity')
def security_identity(asset_id:int,x:IdentityIn,request:Request):
    require_role(request,'Admin','Operator')
    if not x.authorized or not (x.observed_mac or x.observed_hostname):
        raise HTTPException(400,'Provide an actual observed MAC/hostname and confirm its source; baseline is not evidence')
    with connection() as c:
        a=c.execute('SELECT ip FROM security_assets WHERE id=?',(asset_id,)).fetchone()
    if not a: raise HTTPException(404,'Asset not found')
    from modules.enterprise_security import verify_asset_identity
    result=verify_asset_identity(a['ip'],x.observed_mac,x.observed_hostname)
    return {**result,'observation_source':'manual operator-supplied observation','observed_at':utcnow()}


@router.get('/history/{device_id}')
def history(device_id:int,request:Request,days:int=1,ifindex:int|None=None):
    require_role(request)
    if days not in (1,7,30): raise HTTPException(400,'days must be 1, 7 or 30')
    from datetime import timedelta
    # Desktop timestamps use local wall time; convert UTC cutoff to configured local time.
    from zoneinfo import ZoneInfo
    cutoff=(datetime.now(ZoneInfo(os.environ.get('NA_TIMEZONE','Asia/Ho_Chi_Minh')))-timedelta(days=days)).strftime('%Y-%m-%d %H:%M:%S')
    with connection() as c:
        d=c.execute('SELECT * FROM network_devices WHERE id=?',(device_id,)).fetchone()
        if not d: raise HTTPException(404,'Managed device not found')
        host=ip_value(dict(d))
        health=[dict(r) for r in c.execute('SELECT * FROM health_samples WHERE host=? AND created_at>=? ORDER BY id DESC LIMIT 1000',(host,cutoff))]
        indices=[r[0] for r in c.execute('SELECT DISTINCT ifindex FROM interface_samples WHERE host=? ORDER BY ifindex',(host,))]
        interfaces=[dict(r) for r in c.execute('SELECT * FROM interface_samples WHERE host=? AND created_at>=? AND (? IS NULL OR ifindex=?) ORDER BY id DESC LIMIT 1000',(host,cutoff,ifindex,ifindex))]
    return {'host':host,'days':days,'health':health,'interfaces':interfaces,'ifindices':indices,'limit':1000,
            'note':'Up to 1000 latest recorded samples; no interpolation of missing observations.'}
